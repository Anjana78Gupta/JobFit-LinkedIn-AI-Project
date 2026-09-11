"""DOCX -> raw text extraction, using python-docx."""

from __future__ import annotations

from pathlib import Path

from config.logging_config import get_logger
from src.core.exceptions import DocumentExtractionError
from src.parsing.base_parser import AbstractDocumentParser

logger = get_logger(__name__)


class DocxExtractor(AbstractDocumentParser):
    """Extracts raw text from .docx resumes, including table cell contents."""

    @classmethod
    def supported_extensions(cls) -> tuple[str, ...]:
        return (".docx",)

    def extract_text(self, file_path: Path) -> str:
        try:
            import docx
        except ImportError as exc:  # pragma: no cover - dependency guard
            raise DocumentExtractionError(
                "python-docx is not installed. Run `pip install python-docx`.",
            ) from exc

        if not file_path.exists():
            raise DocumentExtractionError(f"File not found: {file_path}")

        try:
            document = docx.Document(str(file_path))
            paragraphs = [p.text for p in document.paragraphs if p.text.strip()]

            table_cells: list[str] = []
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            table_cells.append(cell.text)

            text = "\n".join([*paragraphs, *table_cells]).strip()
        except Exception as exc:  # noqa: BLE001
            raise DocumentExtractionError(
                f"Failed to extract text from DOCX: {file_path.name}",
                details={"error": str(exc)},
            ) from exc

        if not text:
            raise DocumentExtractionError(f"No extractable text found in DOCX: {file_path.name}")

        logger.info("DOCX text extracted", extra={"component": f"chars={len(text)}"})
        return text
